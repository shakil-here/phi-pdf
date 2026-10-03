"""Every tool on small generated PDFs (no network). Bengali checks use digits and words without conjuncts: text
extracted from PDFs garbles conjuncts."""
from __future__ import annotations

import io
import os
import shutil
import zipfile
from pathlib import Path

import pymupdf
import pytest

from pdfstudio import capabilities, ocr
from pdfstudio import tools as T
from pdfstudio.edit_text import page_blocks


def make_pdf(path: Path, pages: int = 6, words: str = "Sample Notes", size=(595.28, 841.89), image: bool = False) -> T.Input:
    doc = pymupdf.open()
    picture = b""
    if image:                                              # one picture, shown on every page
        from PIL import Image
        buffer = io.BytesIO()
        Image.effect_noise((1600, 1600), 60).convert("RGB").save(buffer, "PNG")
        picture = buffer.getvalue()
    xref = 0
    for i in range(pages):
        page = doc.new_page(width=size[0], height=size[1])
        page.insert_text((72, 100), f"{words} page {i + 1}", fontsize=14)
        page.insert_text((72, 140), "Year 2025 exam", fontsize=12)
        if picture:
            rect = pymupdf.Rect(72, 200, 520, 650)
            xref = page.insert_image(rect, xref=xref) if xref else page.insert_image(rect, stream=picture)
    doc.set_toc([[1, "Chapter one", 1], [1, "Chapter two", 3], [2, "Topic", 4], [1, "Chapter three", 5]][: max(1, pages)]
                if pages >= 5 else [])
    doc.save(path)
    return T.Input(path, path.name)


def first(result: T.Result) -> pymupdf.Document:
    return pymupdf.open(result.outputs[0].path)


def test_page_lists_read_ranges_words_and_bengali_digits():
    assert T.page_list("1-3, 7, 10-last", 12) == [0, 1, 2, 6, 9, 10, 11]
    assert T.page_list("৩-১", 5) == [2, 1, 0]
    assert T.page_list("odd", 5) == [0, 2, 4]
    assert T.page_list("2 - 4", 5) == [1, 2, 3]
    assert T.page_list("", 3) == [0, 1, 2]
    assert T.page_list("-2", 5) == [0, 1]
    for bad in ("0", "13", "x", "3-99"):
        with pytest.raises(T.ToolError):
            T.page_list(bad, 12)
    assert T.format_ranges({0, 2, 3, 4}) == "1, 3-5"


def test_merge_split_organize_and_page_tools(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 6)
    b = make_pdf(tmp_path / "b.pdf", 3, words="Second")
    merged = T.run("merge", [a, b], {"odd_blank": True}, tmp_path / "m")
    with first(merged) as doc:
        assert doc.page_count == 9                         # 6 + 3, no blank after an even file
        toc = doc.get_toc()
        assert toc[0][:2] == [1, "a"] and [1, "b", 7] in [t[:3] for t in toc]
    merged = T.run("merge", [b, a], {"odd_blank": True, "ranges": ["1-2", ""]}, tmp_path / "m2")
    with first(merged) as doc:
        assert doc.page_count == 2 + 6                     # only pages 1-2 of b; 2 is even, no blank

    parts = T.run("split", [a], {"mode": "ranges", "ranges": "1-2, 3-last"}, tmp_path / "s")
    assert [pymupdf.open(o.path).page_count for o in parts.outputs] == [2, 4]
    every = T.run("split", [a], {"mode": "every", "every": 4}, tmp_path / "s2")
    assert [pymupdf.open(o.path).page_count for o in every.outputs] == [4, 2]
    chapters = T.run("split", [a], {"mode": "bookmarks", "level": 1}, tmp_path / "s3")
    assert [pymupdf.open(o.path).page_count for o in chapters.outputs] == [2, 2, 2]
    assert "Chapter two" in chapters.outputs[1].name

    plan = [{"f": 1, "p": 0}, {"f": 0, "p": 5, "r": 90}, {"blank": True, "w": 300, "h": 400}, {"f": 0, "p": 5}]
    with first(T.run("organize", [a, b], {"plan": plan}, tmp_path / "o")) as doc:
        assert doc.page_count == 4
        assert "Second page 1" in doc[0].get_text()
        assert doc[1].rotation == 90 and "page 6" in doc[1].get_text()
        assert not doc[2].get_text().strip() and round(doc[2].rect.width) == 300
        assert "page 6" in doc[3].get_text()

    with first(T.run("remove", [a], {"pages": "2-3"}, tmp_path / "r")) as doc:
        assert doc.page_count == 4 and "page 4" in doc[1].get_text()
    with pytest.raises(T.ToolError):
        T.run("remove", [a], {"pages": "1-6"}, tmp_path / "r2")
    with first(T.run("extract", [a], {"pages": "5, 1"}, tmp_path / "e")) as doc:
        assert [p.get_text().split("\n")[0] for p in doc] == ["Sample Notes page 5", "Sample Notes page 1"]
    with first(T.run("rotate", [a], {"pages": "even", "angle": 270}, tmp_path / "ro")) as doc:
        assert [p.rotation for p in doc] == [0, 270, 0, 270, 0, 270]
    with first(T.run("insert", [a], {"what": "blank", "where": "every", "step": 2, "count": 1}, tmp_path / "i")) as doc:
        assert doc.page_count == 9 and not doc[2].get_text().strip()
    with first(T.run("insert", [a], {"what": "file", "where": "after", "after": 1, "source_pages": "2"}, tmp_path / "i2",
                     extras={"source_file": b})) as doc:
        assert doc.page_count == 7 and "Second page 2" in doc[1].get_text()


def test_compress_makes_pictures_smaller_and_refuses_when_it_cannot(tmp_path):
    big = make_pdf(tmp_path / "big.pdf", 2, image=True)
    result = T.run("compress", [big], {"level": "high"}, tmp_path / "c")
    assert result.stats["after"] < result.stats["before"] * 0.8
    with first(result) as doc:
        assert doc.page_count == 2 and "page 2" in doc[1].get_text()
    with pytest.raises(T.ToolError, match="already as small"):
        T.run("compress", [T.Input(result.outputs[0].path, "x.pdf")], {"level": "low"}, tmp_path / "c2")


def test_stamps_numbers_watermark_crop_resize_and_nup(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 6)
    with first(T.run("numbers", [a], {"position": "br", "format": "n_of", "digits": "bengali", "first": 2,
                                      "header": "Physics Notes HSC", "footer": "example.com"}, tmp_path / "n")) as doc:
        assert "১ / ৫" not in doc[0].get_text() and "১ / ৫" in doc[1].get_text()   # numbering starts on page 2
        assert "Physics Notes HSC" in doc[0].get_text() and "example.com" in doc[5].get_text()
    with first(T.run("watermark", [a], {"text": "CONFIDENTIAL", "angle": 45, "skip": "1"}, tmp_path / "w")) as doc:
        assert "CONFIDENTIAL" not in doc[0].get_text() and "CONFIDENTIAL" in doc[1].get_text()
    logo = tmp_path / "logo.png"
    from PIL import Image
    Image.new("RGBA", (200, 100), (20, 40, 200, 255)).save(logo)
    with first(T.run("watermark", [a], {"kind": "image", "position": "tiled", "rows": 2, "cols": 2}, tmp_path / "w2",
                     extras={"image_file": T.Input(logo, "logo.png", "image")})) as doc:
        assert len(doc[0].get_text("dict")["blocks"]) >= 2
    with first(T.run("crop", [a], {"mode": "margins", "top": 20, "left": 10}, tmp_path / "cr")) as doc:
        assert round(doc[0].rect.height) == round(841.89 - 20 * T.MM)
        assert round(doc[0].rect.width) == round(595.28 - 10 * T.MM)
    with first(T.run("crop", [a], {"mode": "auto", "pad": 2}, tmp_path / "cr2")) as doc:
        assert doc[0].rect.height < 200                    # only the two text lines are left
    with first(T.run("resize", [a], {"size": "A5"}, tmp_path / "rs")) as doc:
        assert round(doc[0].rect.width) == 420 and "page 1" in doc[0].get_text()
    with first(T.run("nup", [a], {"per_sheet": 4}, tmp_path / "nu")) as doc:
        assert doc.page_count == 2
    booklet = T.run("nup", [a], {"mode": "booklet"}, tmp_path / "bk")
    with first(booklet) as doc:                            # 6 pages + 2 blanks = 2 sheets, 4 sides
        assert doc.page_count == 4 and doc[0].rect.width > doc[0].rect.height
        text = doc[0].get_text()
        assert "page 1" in text and "page 2" not in text   # front of sheet 1: blank (8) + page 1
    png = T.preview("watermark", [a], {"text": "DRAFT"}, 2, 400)
    assert png.startswith(b"\x89PNG")
    assert T.preview("nup", [a], {"per_sheet": 2}, 1, 300).startswith(b"\x89PNG")


def test_convert_text_pictures_and_images_to_pdf(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 3, image=True)
    txt = T.run("pdf2text", [a], {}, tmp_path / "t")
    body = txt.outputs[0].path.read_text(encoding="utf-8")
    assert "===== Page 3 =====" in body and "page 3" in body
    pics = T.run("pdf2img", [a], {"pages": "1-2", "format": "jpg", "dpi": 50}, tmp_path / "p")
    assert len(pics.outputs) == 2 and pics.outputs[0].path.read_bytes()[:3] == b"\xff\xd8\xff"
    found = T.run("extract_images", [a], {}, tmp_path / "x")
    assert len(found.outputs) == 1                         # the same picture on 3 pages is saved once

    from PIL import Image
    photo = tmp_path / "photo.jpg"
    img = Image.new("RGB", (400, 300), "orange")
    exif = img.getexif()
    exif[0x0112] = 6                                       # a phone photo stored sideways
    img.save(photo, exif=exif)
    png = tmp_path / "scan.png"
    Image.new("L", (300, 300), 200).save(png)
    with first(T.run("img2pdf", [T.Input(photo, "photo.jpg", "image"), T.Input(png, "scan.png", "image")],
                     {"size": "fit", "margin": "none"}, tmp_path / "ip")) as doc:
        assert doc.page_count == 2
        assert doc[0].rect.height > doc[0].rect.width      # turned upright: 300 × 400


def test_protect_unlock_metadata_flatten(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 2)
    locked = T.run("protect", [a], {"password": "secret1", "allow_print": True}, tmp_path / "p")
    with first(locked) as doc:
        assert doc.needs_pass
        assert doc.authenticate("secret1")
    lock_in = T.Input(locked.outputs[0].path, "locked.pdf")
    with pytest.raises(T.ToolError, match="Unlock"):
        T.run("rotate", [lock_in], {"angle": 90}, tmp_path / "r")
    with pytest.raises(T.ToolError, match="Wrong password"):
        T.run("unlock", [lock_in], {"password": "nope"}, tmp_path / "u0")
    with first(T.run("unlock", [lock_in], {"password": "secret1"}, tmp_path / "u")) as doc:
        assert not doc.needs_pass and not doc.is_encrypted and doc.page_count == 2
    with first(T.run("metadata", [a], {"title": "পদার্থবিজ্ঞান", "author": "PHI PDF Studio"}, tmp_path / "m")) as doc:
        assert doc.metadata["title"] == "পদার্থবিজ্ঞান" and doc.metadata["author"] == "PHI PDF Studio"
    doc = pymupdf.open(a.path)
    doc[0].add_text_annot((100, 100), "note")
    doc.save(tmp_path / "annot.pdf")
    with first(T.run("flatten", [T.Input(tmp_path / "annot.pdf", "annot.pdf")], {}, tmp_path / "f")) as out:
        assert not list(out[0].annots())


def test_edit_redact_and_replace_change_the_text(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 2)
    with pymupdf.open(a.path) as doc:
        hit = doc[0].search_for("2025")[0]
    ops = [{"page": 0, "type": "text", "x": 72, "y": 300, "w": 300, "h": 40, "text": "নতুন লেখা New text ১২৩", "size": 16},
           {"page": 0, "type": "redact", "x": hit.x0 - 1, "y": hit.y0 - 1, "w": hit.width + 2, "h": hit.height + 2},
           {"page": 1, "type": "rect", "x": 50, "y": 50, "w": 100, "h": 50, "color": "#ff0000"},
           {"page": 1, "type": "arrow", "x1": 10, "y1": 10, "x2": 200, "y2": 200},
           {"page": 1, "type": "pen", "points": [[10, 10], [20, 30], [40, 10]]},
           {"page": 1, "type": "image", "x": 300, "y": 300, "w": 100, "h": 100,
            "data": "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="}]
    with first(T.run("edit", [a], {"ops": ops}, tmp_path / "e")) as doc:
        text = doc[0].get_text()
        assert "New text" in text and "১২৩" in text and "2025" not in text and "Year" in text
        assert len(doc[1].get_drawings()) >= 3 and doc[1].get_images()
    with pytest.raises(T.ToolError):
        T.run("edit", [a], {"ops": []}, tmp_path / "e2")
    with first(T.run("replace", [a], {"find": "2025", "replace": "2026"}, tmp_path / "rp")) as doc:
        assert "2026" in doc[0].get_text() and "2025" not in doc[1].get_text()
    with pytest.raises(T.ToolError, match="not found"):
        T.run("replace", [a], {"find": "Chemistry", "replace": "x"}, tmp_path / "rp2")


def test_tools_check_their_inputs(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 1)
    with pytest.raises(T.ToolError, match="2 files"):
        T.run("merge", [a], {}, tmp_path / "x")
    with pytest.raises(T.ToolError, match="picture"):
        T.run("rotate", [T.Input(tmp_path / "p.png", "p.png", "image")], {}, tmp_path / "x")
    assert T.kind_of("a.pdf", b"%PDF-1.7") == "pdf"
    assert T.kind_of("a.png", b"\x89PNG\r\n") == "image"
    assert T.kind_of("a.exe", b"MZ") is None
    assert T.kind_of("a.png", b"%PDF") == "pdf"           # the content decides, not the name




# ---------------------------------------------------------------------------------------------- new in PHI PDF Studio
def _blocks(path: Path, page: int = 0) -> list[dict]:
    with pymupdf.open(path) as doc:
        return page_blocks(doc, page)["blocks"]


def _text_op(page: int, b: dict, text: str, **extra) -> dict:
    x0, y0, x1, y1 = b["bbox"]
    lines = [{"x": ln["bbox"][0], "y": ln["bbox"][1], "w": ln["bbox"][2] - ln["bbox"][0], "h": ln["bbox"][3] - ln["bbox"][1],
              "ox": ln["origin"][0], "oy": ln["origin"][1], "size": ln["size"]} for ln in b["lines"]]
    return {"type": "textedit", "page": page, "x": x0, "y": y0, "w": x1 - x0, "h": y1 - y0, "lines": lines, "text": text,
            "orig": b["text"], "size": b["size"], "old_size": b["size"], "color": b["color"], "font": "original",
            "fontname": b["font"], "bold": b["bold"], "italic": b["italic"], "align": b["align"], **extra}


def test_edit_text_in_place_keeps_everything_else(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect((60, 60, 500, 220), color=(0, 0, 1), fill=(0.9, 0.95, 1))
    page.insert_text((72, 100), "Invoice for March 2025", fontsize=18, fontname="hebo")
    page.insert_text((72, 130), "Total amount: 1,200", fontsize=12, fontname="tiro")
    page.insert_text((72, 146), "Paid by cash.", fontsize=12, fontname="tiro")
    page.insert_text((72, 300), "Footer stays.", fontsize=10)
    turned = doc.new_page()
    turned.insert_text((100, 700), "Turned page title", fontsize=16, rotate=90)
    turned.set_rotation(90)                                # reads horizontally on screen
    doc.save(tmp_path / "a.pdf")
    blocks = _blocks(tmp_path / "a.pdf")
    title = next(b for b in blocks if b["text"].startswith("Invoice"))
    body = next(b for b in blocks if b["text"].startswith("Total"))
    assert title["bold"] and title["family"] == "sans" and body["family"] == "serif" and len(body["lines"]) == 2
    turned_block = _blocks(tmp_path / "a.pdf", 1)[0]
    assert turned_block["text"] == "Turned page title" and turned_block["bbox"][2] > turned_block["bbox"][0] + 80

    ops = [_text_op(0, title, "Invoice for April 2026"), _text_op(0, body, "Total amount: 2,450\nPaid by bKash.\nThank you!"),
           _text_op(1, turned_block, "Turned page (edited)")]
    with pymupdf.open(T.run("edit", [T.Input(tmp_path / "a.pdf", "a.pdf")], {"ops": ops}, tmp_path / "o").outputs[0].path) as out:
        text = out[0].get_text()
        assert "April 2026" in text and "March" not in text and "2,450" in text and "Thank you!" in text
        assert "Footer stays." in text and "1,200" not in text and "cash" not in text
        assert len(out[0].get_drawings()) == 1                 # the blue box is untouched
        spans = [s for b in out[0].get_text("dict")["blocks"] for ln in b.get("lines", []) for s in ln["spans"]
                 if abs(s["origin"][1] - 100) < 0.5 and s["text"].strip()]        # all on the same baseline as before
        assert " ".join(s["text"].strip() for s in sorted(spans, key=lambda s: s["origin"][0])) == "Invoice for April 2026"
        assert all(s["flags"] & 16 or "Bold" in s["font"] for s in spans)
        assert "Turned page (edited)" in out[1].get_text()

    bengali = [_text_op(0, title, "চালান ২০২৬")]
    with pymupdf.open(T.run("edit", [T.Input(tmp_path / "a.pdf", "a.pdf")], {"ops": bengali}, tmp_path / "b").outputs[0].path) as out:
        assert "২০২৬" in out[0].get_text() and "Invoice" not in out[0].get_text()
    emptied = [_text_op(0, body, "")]
    with pymupdf.open(T.run("edit", [T.Input(tmp_path / "a.pdf", "a.pdf")], {"ops": emptied}, tmp_path / "c").outputs[0].path) as out:
        assert "Total" not in out[0].get_text() and "Invoice for March" in out[0].get_text()
    with pytest.raises(T.ToolError, match="Nothing was changed"):   # unchanged text is no edit
        T.run("edit", [T.Input(tmp_path / "a.pdf", "a.pdf")], {"ops": [_text_op(0, title, title["text"])]}, tmp_path / "d")


def _table_pdf(path: Path) -> T.Input:
    doc = pymupdf.open()
    for n in range(2):
        page = doc.new_page()
        page.insert_text((72, 90), f"Report {n + 1}: mail a.b@example.com or call +8801712345678", fontsize=12)
        page.insert_text((72, 110), "Nothing secret here.", fontsize=12)
        for r in range(4):
            for c in range(3):
                cell = pymupdf.Rect(72 + c * 100, 200 + r * 20, 172 + c * 100, 220 + r * 20)
                page.draw_rect(cell, width=0.5)
                page.insert_text((cell.x0 + 4, cell.y1 - 6), "Name" if r == 0 and c == 0 else f"{r * 10 + c}", fontsize=10)
    doc.save(path)
    return T.Input(path, path.name)


def test_pdf_to_word_excel_and_powerpoint(tmp_path):
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    a = _table_pdf(tmp_path / "report.pdf")
    for mode in ("layout", "text"):
        result = T.run("pdf2word", [a], {"mode": mode}, tmp_path / f"w-{mode}")
        assert result.outputs[0].name == "report.docx"
        text = "\n".join(p.text for p in Document(str(result.outputs[0].path)).paragraphs)
        tables = "\n".join(c.text for t in Document(str(result.outputs[0].path)).tables for r in t.rows for c in r.cells)
        assert "Report 1" in text and ("21" in text or "21" in tables)
    book = load_workbook(T.run("pdf2excel", [a], {}, tmp_path / "x").outputs[0].path)
    assert len(book.sheetnames) == 2
    sheet = book[book.sheetnames[0]]
    assert sheet["A1"].value == "Name" and sheet["B2"].value == 11 and sheet["A1"].font.bold
    deck = Presentation(str(T.run("pdf2ppt", [a], {}, tmp_path / "p").outputs[0].path))
    assert len(deck.slides) == 2 and "Report 2" in deck.slides[1].notes_slide.notes_text_frame.text

    scan = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), pixmap=pymupdf.Pixmap(pymupdf.csRGB, (0, 0, 60, 80), 0))
    doc.save(scan)
    with pytest.raises(T.ToolError, match="OCR"):
        T.run("pdf2word", [T.Input(scan, "scan.pdf")], {}, tmp_path / "s")


def test_text_markdown_and_html_to_pdf(tmp_path):
    (tmp_path / "notes.md").write_text("# Heading\n\nA paragraph with **bold**.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n",
                                       encoding="utf-8")
    (tmp_path / "page.html").write_text("<html><body><h2>Web page</h2><script>alert(1)</script>"
                                        "<img src='http://example.com/x.png'><p>বাংলা ১২৩</p></body></html>",
                                        encoding="utf-8")
    inputs = [T.Input(tmp_path / "notes.md", "notes.md", "text"), T.Input(tmp_path / "page.html", "page.html", "text")]
    result = T.run("text2pdf", inputs, {"paper": "A5"}, tmp_path / "t")
    with pymupdf.open(result.outputs[0].path) as doc:
        text = "".join(p.get_text() for p in doc)
        assert doc.page_count == 2 and round(doc[0].rect.width) == 420
        assert "Heading" in text and "Web page" in text and "১২৩" in text and "alert" not in text
        assert [t[1] for t in doc.get_toc()] == ["notes", "page"]
    separate = T.run("text2pdf", inputs, {"merge": False}, tmp_path / "t2")
    assert len(separate.outputs) == 2


def test_auto_redact_and_clean_hidden_data(tmp_path):
    a = _table_pdf(tmp_path / "r.pdf")
    marked = T.run("redact", [a], {"find": ["email"], "mode": "mark"}, tmp_path / "m")
    with pymupdf.open(marked.outputs[0].path) as doc:
        assert "a.b@example.com" in doc[0].get_text()           # only outlined, for checking
    done = T.run("redact", [a], {"find": ["email", "phone"], "words": "secret"}, tmp_path / "r")
    assert done.stats["count"] == 6
    with pymupdf.open(done.outputs[0].path) as doc:
        text = doc[0].get_text()
        assert "example.com" not in text and "8801712345678" not in text and "secret" not in text
        assert "Report 1" in text and "Nothing" in text
    with pytest.raises(T.ToolError, match="Nothing matching"):
        T.run("redact", [a], {"words": "Chemistry"}, tmp_path / "r2")
    with pytest.raises(T.ToolError, match="pattern is not valid"):
        T.run("redact", [a], {"pattern": "(["}, tmp_path / "r3")

    doc = pymupdf.open(a.path)
    doc.set_metadata({"title": "Secret plan", "author": "Someone"})
    doc.embfile_add("notes.txt", b"hidden attachment")
    doc[0].add_text_annot((100, 100), "a private note")
    doc.save(tmp_path / "dirty.pdf")
    clean = T.run("sanitize", [T.Input(tmp_path / "dirty.pdf", "dirty.pdf")], {"comments": True}, tmp_path / "s")
    with pymupdf.open(clean.outputs[0].path) as out:
        assert out.embfile_count() == 0 and not out.metadata.get("author") and not list(out[0].annots())
        assert "Report 1" in out[0].get_text()


@pytest.mark.skipif(not capabilities.tesseract_path(), reason="Tesseract is not installed")
def test_ocr_reads_a_tilted_scan_into_pdf_text_and_word(tmp_path):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 64)
    except OSError:
        font = ImageFont.load_default(size=64)
    draw.text((150, 300), "Searchable scanned page", fill="black", font=font)
    draw.text((150, 420), "Physics question number 42", fill="black", font=font)
    tilted = image.rotate(2.0, fillcolor="white", resample=Image.BICUBIC)
    assert abs(ocr.skew_angle(tilted) + 2.0) < 0.35          # the turn that levels it: back by 2°
    buf = io.BytesIO()
    tilted.save(buf, "JPEG", quality=70)                       # JPEG noise in the white paper, like real scans
    doc = pymupdf.open()
    doc.new_page(width=612, height=792).insert_image(pymupdf.Rect(0, 0, 612, 792), stream=buf.getvalue())
    doc.save(tmp_path / "scan.pdf")
    # a language "downloaded" into the app's own folder: Tesseract then runs with that folder, which must also get
    # the system's output settings (configs/pdf), or the searchable PDF comes out without text
    own = capabilities.Tessdata(tmp_path / "tessdata")
    own.own.mkdir()
    shutil.copyfile(own.system_dir() / "eng.traineddata", own.own / "eng.traineddata")
    ocr.TESSDATA = own
    assert own.folder() == own.own and (own.own / "configs" / "pdf").exists()
    scan = T.Input(tmp_path / "scan.pdf", "scan.pdf")
    result = T.run("ocr", [scan], {"language": ["eng"], "output": "pdf_txt", "dpi": 200}, tmp_path / "o")
    pdf, txt = result.outputs
    with pymupdf.open(pdf.path) as out:
        assert "question number 42" in out[0].get_text().lower()
    assert "searchable" in txt.path.read_text(encoding="utf-8").lower()
    assert "straightened" in result.summary
    with pytest.raises(T.ToolError, match="already has text"):
        T.run("ocr", [T.Input(pdf.path, "x.pdf")], {"language": ["eng"]}, tmp_path / "again")
    word = T.run("ocr", [scan], {"language": ["eng"], "output": "docx", "clean": "strong"}, tmp_path / "w")
    from docx import Document
    assert "42" in "\n".join(p.text for p in Document(str(word.outputs[0].path)).paragraphs)
    with pytest.raises(T.ToolError, match="Not installed for OCR"):
        T.run("ocr", [scan], {"language": ["zzz"]}, tmp_path / "x")


def test_ai_reading_needs_a_key_and_reports_provider_errors(tmp_path, monkeypatch):
    from pdfstudio import ai

    a = _table_pdf(tmp_path / "a.pdf")
    monkeypatch.setattr(ai, "USE_SERVER_KEYS", False)
    with pytest.raises(T.ToolError, match="API key"):
        T.run("ocr", [a], {"engine": "ai", "ai_provider": "gemini"}, tmp_path / "x")
    seen = []

    def fake(png, provider, model, key, extra=""):
        seen.append((provider, model, key, png[:4]))
        return "# Page\n\n$E = mc^2$ বাংলা"
    monkeypatch.setattr(ai, "transcribe", fake)
    result = T.run("ocr", [a], {"engine": "ai", "ai_provider": "claude", "ai_key": "sk-test", "output": "md"}, tmp_path / "y")
    body = result.outputs[0].path.read_text(encoding="utf-8")
    assert body.count("$E = mc^2$") == 2 and "<!-- page 2 -->" in body
    assert seen[0][:3] == ("claude", "claude-opus-5-5", "sk-test") and seen[0][3] == b"\x89PNG"


def test_new_file_kinds_are_recognised():
    assert T.kind_of("a.docx", b"PK\x03\x04rest") == "office"
    assert T.kind_of("a.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") == "office"
    assert T.kind_of("a.docx", b"%PDF-1.7") == "pdf"                 # the content decides
    assert T.kind_of("a.docx", b"MZ\x90") is None
    assert T.kind_of("notes.md", "# শিরোনাম".encode()) == "text"
    assert T.kind_of("page.html", b"<html>") == "text"
    assert T.kind_of("x.txt", b"\x00\x01binary") is None
    assert T.kind_of("x.svg", b"<svg>") is None                       # never accepted (scripts inside)


@pytest.mark.skipif(not capabilities.libreoffice_path(), reason="LibreOffice is not installed")
def test_office_to_pdf(tmp_path):
    from docx import Document
    document = Document()
    document.add_heading("Converted with LibreOffice", 1)
    document.save(tmp_path / "a.docx")
    result = T.run("office2pdf", [T.Input(tmp_path / "a.docx", "a.docx", "office")], {}, tmp_path / "o")
    with pymupdf.open(result.outputs[0].path) as doc:
        assert "Converted with LibreOffice" in doc[0].get_text()


def test_office_to_pdf_says_how_to_install_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr(capabilities, "libreoffice_path", lambda: None)
    (tmp_path / "a.docx").write_bytes(b"PK\x03\x04")
    with pytest.raises(T.ToolError, match="LibreOffice"):
        T.run("office2pdf", [T.Input(tmp_path / "a.docx", "a.docx", "office")], {}, tmp_path / "o")


# ---------------------------------------------------------------------------------------------- editor: more tools
def test_editor_marks_shapes_stamps_notes_and_links(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 1)
    box = {"x": 300, "y": 300, "w": 80, "h": 60, "color": "#2563eb", "width": 2}
    ops = [{"page": 0, "type": kind, **box, "dash": dash} for kind, dash in
           (("roundrect", "dash"), ("triangle", "dot"), ("diamond", "solid"), ("star", "solid"))]
    ops += [{"page": 0, "type": kind, "x": 72, "y": 400, "w": 20, "h": 20, "color": "#16a34a"} for kind in ("check", "cross", "dot")]
    ops += [{"page": 0, "type": kind, "points": [[80, 500], [120, 520], [160, 500]], "width": 4} for kind in
            ("pencil", "marker", "whitener")]
    ops += [{"page": 0, "type": "arrow2", "x1": 80, "y1": 600, "x2": 300, "y2": 620},
            {"page": 0, "type": "underline", "x": 72, "y": 88, "w": 120, "h": 16},
            {"page": 0, "type": "stamp", "x": 350, "y": 80, "w": 150, "h": 44, "text": "PAID", "color": "#dc2626"},
            {"page": 0, "type": "note", "x": 500, "y": 200, "w": 160, "h": 60, "text": "Check the total"},
            {"page": 0, "type": "link", "x": 72, "y": 700, "w": 100, "h": 20, "url": "example.org/help"},
            {"page": 0, "type": "text", "x": 72, "y": 740, "w": 300, "h": 30, "text": "Signed by hand", "font": "greatvibes",
             "strike": True}]
    with first(T.run("edit", [a], {"ops": ops}, tmp_path / "e")) as doc:
        page = doc[0]
        assert "PAID" in page.get_text() and "Signed by hand" in page.get_text()
        assert any("Great Vibes" in f[3] for f in page.get_fonts())
        assert [x.info["content"] for x in page.annots()] == ["Check the total"]
        assert page.get_links()[0]["uri"] == "https://example.org/help"
        assert len(page.get_drawings()) >= 14
    with pytest.raises(T.ToolError, match="web address"):
        T.run("edit", [a], {"ops": [{"page": 0, "type": "link", "x": 1, "y": 1, "w": 9, "h": 9, "url": "javascript:x"}]},
              tmp_path / "bad")


def test_font_catalog_has_bundled_fonts_for_bengali_and_handwriting():
    from pdfstudio import fonts as F
    cat = F.catalog()
    assert {"sans", "serif", "mono", "roboto", "hindsiliguri", "notoserifbengali", "dancingscript"} <= set(cat)
    assert cat["roboto"].face(True, True).name == "Roboto-BoldItalic.ttf"
    assert cat["greatvibes"].face(True, False).name == "GreatVibes-Regular.ttf"     # nearest face there is
    assert "PHI Hind Siliguri" in F.stylesheet() and F.bundled_file("../x.ttf") is None
    assert F.bundled_file("Roboto-Regular.ttf") is not None


def test_edit_text_in_another_font(tmp_path):
    a = make_pdf(tmp_path / "a.pdf", 1)
    title = _blocks(a.path)[0]
    op = _text_op(0, title, "Renamed in Roboto", font="roboto", bold=True)
    with first(T.run("edit", [a], {"ops": [op]}, tmp_path / "e")) as doc:
        assert "Renamed in Roboto" in doc[0].get_text() and "Sample Notes page 1" not in doc[0].get_text()
        assert any("Roboto" in f[3] and "Bold" in f[3] for f in doc[0].get_fonts())


def _scan(path: Path, hidden_layer: bool = False) -> T.Input:
    """A page that is only a picture of text (and, if asked, an invisible OCR layer over it)."""
    src = pymupdf.open()
    page = src.new_page()
    page.draw_rect(page.rect, fill=(0.96, 0.94, 0.86), color=None)
    page.insert_text((72, 100), "Invoice number 12345", fontsize=18)
    page.insert_text((72, 160), "Total due today", fontsize=14)
    picture = page.get_pixmap(dpi=200).tobytes("png")
    doc = pymupdf.open()
    out = doc.new_page(width=page.rect.width, height=page.rect.height)
    out.insert_image(out.rect, stream=picture)
    if hidden_layer:
        out.insert_text((72, 100), "Invoice number 12345", fontsize=18, render_mode=3)
    doc.save(path)
    return T.Input(path, path.name)


def test_hidden_ocr_text_is_marked_and_covered_when_changed(tmp_path):
    a = _scan(tmp_path / "s.pdf", hidden_layer=True)
    with pymupdf.open(a.path) as doc:
        data = page_blocks(doc, 0)
    b = data["blocks"][0]
    assert b["ocr"] and b["bg"].startswith("#f")
    op = _text_op(0, b, "Invoice number 99999", ocr=True, bg=b["bg"])
    with first(T.run("edit", [a], {"ops": [op]}, tmp_path / "e")) as doc:
        assert "99999" in doc[0].get_text() and "12345" not in doc[0].get_text()
        pix = doc[0].get_pixmap(clip=pymupdf.Rect(b["bbox"]) + (0, 0, 40, 0), dpi=72)
        assert min(pix.pixel(x, y)[0] for x in range(pix.width) for y in range(pix.height)) < 90   # new letters drawn


@pytest.mark.skipif(not capabilities.tesseract_path() or "eng" not in ocr.TESSDATA.languages(),
                    reason="Tesseract with English is not installed")
def test_ocr_blocks_make_a_scanned_page_editable(tmp_path):
    from pdfstudio.edit_text import ocr_blocks
    a = _scan(tmp_path / "s.pdf")
    with pymupdf.open(a.path) as doc:
        assert page_blocks(doc, 0)["scanned"]
        data = ocr_blocks(doc, 0, "eng")
    texts = [b["text"] for b in data["blocks"]]
    assert "Invoice number 12345" in texts and "Total due today" in texts
    b = data["blocks"][texts.index("Invoice number 12345")]
    assert 15 < b["size"] < 22 and b["ocr"] and b["bg"] != "#ffffff"
    with first(T.run("edit", [a], {"ops": [_text_op(0, b, "Invoice number 77", ocr=True, bg=b["bg"])]}, tmp_path / "e")) as doc:
        assert doc[0].get_text().strip() == "Invoice number 77"


def _dark_slide(path: Path) -> T.Input:
    """Light text on a dark slide, one word in another colour, as a picture."""
    src = pymupdf.open()
    page = src.new_page(width=800, height=300)
    page.draw_rect(page.rect, fill=(0.02, 0.02, 0.03), color=None)
    page.insert_text((60, 100), "Find the major axis of the ellipse", fontsize=24, fontname="tiro", color=(1, 1, 1))
    page.insert_text((540, 100), "[RU 2021]", fontsize=24, fontname="tiro", color=(0.96, 0.62, 0.2))
    page.insert_text((60, 200), "Second line stays as it is", fontsize=20, fontname="helv", color=(1, 1, 1))
    picture = page.get_pixmap(dpi=200).tobytes("png")
    doc = pymupdf.open()
    out = doc.new_page(width=800, height=300)
    out.insert_image(out.rect, stream=picture)
    doc.save(path)
    return T.Input(path, path.name)


@pytest.mark.skipif(not capabilities.tesseract_path() or "eng" not in ocr.TESSDATA.languages(),
                    reason="Tesseract with English is not installed")
def test_ocr_on_a_dark_slide_keeps_its_look_and_changes_only_the_changed_words(tmp_path):
    from pdfstudio.edit_text import ocr_blocks
    a = _dark_slide(tmp_path / "d.pdf")
    with pymupdf.open(a.path) as doc:
        blocks = ocr_blocks(doc, 0, "eng")["blocks"]
    b = next(x for x in blocks if "major" in x["text"])
    ink, paper = (int(b["color"][1:3], 16), int(b["bg"][1:3], 16))
    assert ink > 200 and paper < 40                        # light letters on dark paper, not the other way round
    assert b["runs"] and any(r["c"][1:3] > "c0" and r["c"][3:5] < "c0" for r in b["runs"])   # the orange words
    assert b["fontkey"] == "serif" and 21 < b["size"] < 27
    new = b["text"].replace("major", "minor")
    runs = [dict(r, t=r["t"].replace("major", "minor")) for r in b["runs"]]
    op = _text_op(0, b, new, ocr=True, bg=b["bg"], font=b["fontkey"], runs=runs)
    for ln, src in zip(op["lines"], b["lines"]):
        ln["words"] = src["words"]
    first_word = pymupdf.Rect(b["lines"][0]["words"][0][:4])
    with pymupdf.open(a.path) as doc:
        before = doc[0].get_pixmap(dpi=72, clip=first_word).samples
    with first(T.run("edit", [a], {"ops": [op]}, tmp_path / "e")) as doc:
        assert "minor" in doc[0].get_text() and "major" not in doc[0].get_text()
        assert "Find" not in doc[0].get_text()             # unchanged words were not written again…
        assert doc[0].get_pixmap(dpi=72, clip=first_word).samples == before     # …they are the same pixels
        assert len(doc[0].get_images()) >= 2              # and the words after the change moved as pixels


def test_edit_text_changes_only_the_changed_words(tmp_path):
    from pdfstudio import edit_text as E
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=260)
    lora = pymupdf.Font(fontfile=str(Path(E.__file__).parent / "fonts" / "Lora-Bold.ttf"))
    page.insert_font(fontname="L", fontbuffer=lora.buffer)
    page.insert_text((30, 50), "Admission Test 2025: Result Sheet", fontname="L", fontsize=16)
    page.insert_text((30, 90), "The exam was held on 12 March in Dhaka.", fontsize=12)
    page.insert_text((30, 106), "Candidates must bring their admit card.", fontsize=12)
    page.insert_text((30, 122), "Total marks: 100, pass marks: 40.", fontsize=12, fontname="tiro")
    doc.subset_fonts()                                     # like most PDFs: the font is cut down, its cmap dropped
    doc.save(tmp_path / "w.pdf")
    title, body = _blocks(tmp_path / "w.pdf")[:2]
    with pymupdf.open(tmp_path / "w.pdf") as src:
        fonts = E.PageFonts(src[0])
        cut = fonts.font("Lora Bold", 0, "2052")          # the embedded font, its letters found again from the page
        assert cut is not None and len(cut.buffer) < len(lora.buffer) / 4 and cut.has_glyph(ord("5"))
        assert len(fonts.font("Lora Bold", 0, "Final").buffer) == len(lora.buffer)   # "F": not in it, the full font
        keep = [pymupdf.Rect(30, 36, 150, 54), pymupdf.Rect(30, 78, 150, 94), pymupdf.Rect(30, 110, 300, 126)]
        before = [src[0].get_pixmap(dpi=96, clip=r).samples for r in keep]

    ops = [_text_op(0, title, "Admission Test 2052: Final Result Sheet"),
           _text_op(0, body, body["text"].replace("12 March", "5 April").replace("their admit", "the"))]
    with first(T.run("edit", [T.Input(tmp_path / "w.pdf", "w.pdf")], {"ops": ops}, tmp_path / "e")) as out:
        text = out[0].get_text()
        for gone in ("2025", "12 March", "their", "admit"):
            assert gone not in text
        words = [w[4] for w in out[0].get_text("words")]
        for new in ("2052:", "Final", "5", "April", "the", "card.", "Dhaka."):
            assert new in words
        # words before a change, and lines nobody changed, are the very same letters
        assert [out[0].get_pixmap(dpi=96, clip=r).samples for r in keep] == before
        assert "Admission Test" in text and "Total marks: 100, pass marks: 40." in text
        spans = [s for b in out[0].get_text("dict")["blocks"] for ln in b.get("lines", []) for s in ln["spans"]]
        assert all("Lora" in s["font"] for s in spans if s["text"].strip() in ("2052:", "Final", "Result", "Sheet"))
        line = sorted((s for s in spans if abs(s["origin"][1] - 90) < 0.5 and s["text"].strip()), key=lambda s: s["origin"][0])
        gaps = [b["bbox"][0] - a["bbox"][2] for a, b in zip(line, line[1:])]
        assert all(g < 6 for g in gaps)                  # the words after a shorter one moved up to it


@pytest.mark.skipif(os.name == "nt", reason="Windows shortcuts are made with PowerShell")
def test_desktop_icon_and_shortcut(tmp_path, monkeypatch):
    from PIL import Image

    from pdfstudio import desktop
    for name, size in (("i.icns", 1024), ("i.ico", 256), ("i.png", 256)):
        with Image.open(desktop.write_icon(tmp_path / name)) as icon:
            assert icon.size == (size, size) and icon.getpixel((size // 2, size // 2))[:3] == (255, 255, 255)
    monkeypatch.setenv("HOME", str(tmp_path / "home"))
    monkeypatch.setenv("XDG_DATA_HOME", str(tmp_path / "home" / ".local" / "share"))
    monkeypatch.setattr(desktop, "_desktop", lambda: tmp_path / "home" / "Desktop")
    (tmp_path / "home" / "Desktop").mkdir(parents=True)
    made = desktop.install_shortcuts(tmp_path / "data")
    assert any(p.parent == tmp_path / "home" / "Desktop" for p in made)
    runner = next((p / "Contents" / "MacOS" / "phi-pdf-studio" for p in made if p.suffix == ".app"), None) or made[0]
    assert "pdfstudio launch --data" in runner.read_text() and str(tmp_path / "data") in runner.read_text()
    assert desktop.running_url(tmp_path / "data") is None    # nothing running yet
