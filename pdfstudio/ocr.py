"""OCR: read scanned pages, in any of Tesseract's 100+ languages (several at once: "ben+eng"), or with an AI model.

Tesseract runs as its own program, one page per process, several pages at a time. Before reading, a page can be
cleaned (grey, contrast, noise, black-and-white) and straightened (tilt found by testing small angles; sideways
pages found by Tesseract's orientation check). The result can be
  - a searchable PDF: the page keeps its look and gets an invisible text layer (or, when it was straightened,
    the straightened picture);
  - plain text (.txt) or a Word file (.docx), page by page.
The AI engine (ai.py) reads pages with Gemini or Claude for hard pages (Bengali with math, handwriting, complex
tables) and gives Markdown, text or Word; it needs the person's own API key.
"""
from __future__ import annotations

import io
import os
import re
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pymupdf

from . import capabilities as C
from . import tools as T
from .tools import Input, Progress, Result, ToolError

TESSDATA = C.Tessdata(Path.home() / ".phi-pdf-studio" / "tessdata")   # replaced by the app with its data folder
PSM = {"auto": "3", "column": "4", "block": "6", "sparse": "11", "line": "7"}
OUTPUTS = ("pdf", "txt", "docx", "pdf_txt", "md")
PAGE_TIMEOUT = 300


# ---------------------------------------------------------------------------------------------- pictures
def _to_image(pix: pymupdf.Pixmap):
    from PIL import Image
    mode = "L" if pix.n == 1 else "RGB"
    return Image.frombytes(mode, (pix.width, pix.height), pix.samples)


def _otsu(grey) -> int:
    hist = grey.histogram()[:256]
    total = sum(hist)
    sum_all = sum(i * h for i, h in enumerate(hist))
    weight_b = sum_b = 0
    best, threshold = -1.0, 128
    for i, h in enumerate(hist):
        weight_b += h
        if weight_b == 0:
            continue
        weight_f = total - weight_b
        if weight_f == 0:
            break
        sum_b += i * h
        mean_b, mean_f = sum_b / weight_b, (sum_all - sum_b) / weight_f
        between = weight_b * weight_f * (mean_b - mean_f) ** 2
        if between > best:
            best, threshold = between, i
    return threshold


def clean(image, level: str):
    """none | auto (grey + contrast) | strong (+ noise removed, pure black and white)."""
    from PIL import ImageFilter, ImageOps

    if level == "none":
        return image
    # a plain min-max stretch: a cutoff would count JPEG noise in the white paper as "the darkest 1%" when the
    # text itself is less than 1% of the page, and blacken the whole background
    grey = ImageOps.autocontrast(image.convert("L"), cutoff=0)
    if level == "auto":
        return grey
    grey = grey.filter(ImageFilter.MedianFilter(3))
    cut = _otsu(grey)
    return grey.point(lambda v: 255 if v > cut else 0, mode="1").convert("L")


def skew_angle(image) -> float:
    """Degrees to turn the page (counter-clockwise, as PIL's rotate) so its text rows are level (−6…6); 0 when
    it is level already. Found by testing small turns and keeping the one whose rows of ink are sharpest."""
    try:
        import numpy as np
    except ImportError:      # numpy is optional: no straightening without it
        return 0.0
    from PIL import ImageOps

    small = ImageOps.autocontrast(image.convert("L"))
    small.thumbnail((900, 900))
    cut = _otsu(small)
    ink = small.point(lambda v: 255 if v < cut else 0)

    def score(angle: float) -> float:
        rows = np.asarray(ink.rotate(angle, fillcolor=0), dtype=np.float32).sum(axis=1)
        return float(np.sum(np.diff(rows) ** 2))

    best = max((a / 2 for a in range(-12, 13)), key=score)
    fine = max((best + d / 10 for d in range(-5, 6)), key=score)
    return fine if abs(fine) >= 0.2 else 0.0


def _run(cmd: list[str], timeout: float = PAGE_TIMEOUT) -> subprocess.CompletedProcess:
    env = dict(os.environ, OMP_THREAD_LIMIT="1")
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    try:
        return subprocess.run(cmd, capture_output=True, timeout=timeout, env=env, creationflags=flags)
    except subprocess.TimeoutExpired as exc:
        raise ToolError("Reading a page took too long (more than 5 minutes). Try a lower resolution.") from exc


def orientation(binary: str, png: Path, tessdata: Path | None) -> int:
    """Degrees (0, 90, 180, 270) to turn the page clockwise so its text stands upright; 0 when unsure."""
    cmd = [binary, str(png), "stdout", "--psm", "0"] + (["--tessdata-dir", str(tessdata)] if tessdata else [])
    out = _run(cmd, 60)
    text = out.stdout.decode("utf-8", "replace") + out.stderr.decode("utf-8", "replace")
    rotate = re.search(r"Rotate:\s*(\d+)", text)
    conf = re.search(r"Orientation confidence:\s*([\d.]+)", text)
    if rotate and conf and float(conf.group(1)) >= 2.0:
        return int(rotate.group(1)) % 360
    return 0


def read_picture(binary: str, png: bytes, language: str, psm: str, dpi: int, tessdata: Path | None,
                 want_pdf: bool) -> tuple[bytes, str]:
    """Tesseract on one page picture: (a text-only PDF page or b"", the text)."""
    with tempfile.TemporaryDirectory(prefix="phipdf-ocr-") as tmp:
        src = Path(tmp) / "page.png"
        src.write_bytes(png)
        base = Path(tmp) / "out"
        cmd = [binary, str(src), str(base), "-l", language, "--psm", psm, "--dpi", str(dpi)]
        if tessdata:
            cmd += ["--tessdata-dir", str(tessdata)]
        cmd += ["-c", "textonly_pdf=1"] + (["pdf"] if want_pdf else []) + ["txt"]
        out = _run(cmd)
        if out.returncode != 0:
            message = out.stderr.decode("utf-8", "replace").strip().splitlines()
            detail = next((line for line in reversed(message) if line.strip()), "unknown error")
            if "Failed loading language" in "\n".join(message) or "Error opening data file" in "\n".join(message):
                raise ToolError(f"The OCR language “{language}” is not installed. Add it in Settings → OCR languages.")
            raise ToolError(f"Tesseract could not read a page: {detail[:300]}")
        pdf = (base.with_suffix(".pdf")).read_bytes() if want_pdf and base.with_suffix(".pdf").is_file() else b""
        text = base.with_suffix(".txt").read_text(encoding="utf-8", errors="replace") \
            if base.with_suffix(".txt").is_file() else ""
        return pdf, text


# ---------------------------------------------------------------------------------------------- options
def _languages(p: dict, installed: list[str]) -> str:
    asked = p.get("language")
    asked = "+".join(map(str, asked)) if isinstance(asked, list) else str(asked or "")
    wanted = [lang for lang in re.split(r"[+,\s]+", asked[:200]) if lang]
    missing = [lang for lang in wanted if lang not in installed]
    if missing:
        names = ", ".join(C.language_name(m) for m in missing)
        raise ToolError(f"Not installed for OCR: {names}. Add it in Settings → OCR languages.")
    return "+".join(wanted or (["eng"] if "eng" in installed else installed[:1]))


def _docx(pages: list[tuple[int, str]], title: str) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Nirmala UI")   # Bengali/Hindi in Word
    document.core_properties.title = title
    for k, (number, text) in enumerate(pages):
        if k:
            document.add_page_break()
        heading = document.add_paragraph()
        run = heading.add_run(f"Page {number}")
        run.bold = True
        run.font.size = Pt(9)
        for para in re.split(r"\n\s*\n", text.strip()):
            if para.strip():
                document.add_paragraph(re.sub(r"\s*\n\s*", " ", para.strip()))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _text_file(pages: list[tuple[int, str]]) -> str:
    return "".join(f"\n\n===== Page {n} =====\n\n{t.strip()}" for n, t in pages).strip() + "\n"


# ---------------------------------------------------------------------------------------------- the tool
def ocr(inputs: list[Input], p: dict, out_dir: Path, progress: Progress) -> Result:
    engine = T._choice(p, "engine", ("tesseract", "ai"), "tesseract")
    if engine == "ai":
        from .ai import ai_read
        return ai_read(inputs, p, out_dir, progress)
    binary = C.tesseract_path()
    if not binary:
        raise ToolError("Tesseract (the free OCR program) is not installed on this computer. Install it with: "
                        f"{C.INSTALL['tesseract'][C.SYSTEM]} — or use the AI engine.")
    installed = TESSDATA.languages()
    if not installed:
        raise ToolError("Tesseract has no languages installed. Add one in Settings → OCR languages.")
    language = _languages(p, installed)
    tessdata = TESSDATA.folder()
    psm = PSM[T._choice(p, "layout", tuple(PSM), "auto")]
    dpi = T._int(p, "dpi", 300, 150, 600)
    level = T._choice(p, "clean", ("none", "auto", "strong"), "auto")
    straighten = T._bool(p, "deskew", True)
    upright = T._bool(p, "rotate", True) and "osd" in {f.stem for f in (tessdata or TESSDATA.system_dir() or Path("."))
                                                        .glob("*.traineddata")}
    keep = T._choice(p, "keep", ("original", "cleaned"), "original")
    mode = T._choice(p, "mode", ("skip", "all"), "skip")
    output = T._choice(p, "output", OUTPUTS[:4], "pdf")
    inp = inputs[0]
    doc = T.open_pdf(inp)
    pages = T.page_set(T._text(p, "pages", 4000), doc.page_count)
    texts: dict[int, str] = {}
    todo = []
    for i in pages:
        existing = doc[i].get_text("text").strip()
        if mode == "skip" and existing:
            texts[i] = existing
        else:
            todo.append(i)
    if not todo:
        doc.close()
        raise ToolError("Every page already has text, so there was nothing to read. Choose “Read every page” to "
                        "read them again (useful when copied text comes out garbled).")

    def prepare(i: int) -> dict:
        page = T._upright(doc[i])
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)
        image = _to_image(pix)
        turn = 0
        if upright:
            probe = io.BytesIO()
            small = image.copy()
            small.thumbnail((1600, 1600))
            small.save(probe, "PNG")
            with tempfile.TemporaryDirectory(prefix="phipdf-osd-") as tmp:
                path = Path(tmp) / "p.png"
                path.write_bytes(probe.getvalue())
                try:
                    turn = orientation(binary, path, tessdata)
                except ToolError:
                    turn = 0
            if turn:
                image = image.rotate(-turn, expand=True)
        angle = skew_angle(image) if straighten else 0.0
        if angle:
            image = image.rotate(angle, expand=False, fillcolor="white", resample=3)
        cleaned = clean(image, level)
        buf = io.BytesIO()
        cleaned.save(buf, "PNG")
        return {"turn": turn, "angle": angle, "image": image, "cleaned": cleaned, "png": buf.getvalue()}

    want_pdf = output in ("pdf", "pdf_txt")
    workers = max(1, min(4, os.cpu_count() or 2, len(todo)))
    done = 0
    straightened = 0
    progress(0, len(todo), "reading pages")

    def finish(i: int, prep: dict, pdf: bytes, text: str) -> None:
        nonlocal done, straightened
        texts[i] = text
        if want_pdf:
            page = doc[i]                                       # already upright (prepare)
            if prep["turn"] or prep["angle"] or keep == "cleaned":   # the straightened / cleaned picture is the page
                picture = prep["cleaned"] if keep == "cleaned" else prep["image"]
                if prep["turn"] in (90, 270):
                    page.set_mediabox(pymupdf.Rect(0, 0, page.rect.height, page.rect.width))
                for xref in page.get_contents():
                    doc.update_stream(xref, b" ")
                data = io.BytesIO()
                if keep == "cleaned":
                    picture.convert("L").save(data, "PNG", optimize=True)
                else:
                    picture.convert("RGB").save(data, "JPEG", quality=85)
                page.insert_image(page.rect, stream=data.getvalue())
                straightened += bool(prep["turn"] or prep["angle"])
            if pdf:
                layer = pymupdf.open("pdf", pdf)
                page.show_pdf_page(page.rect, layer, 0, overlay=True, keep_proportion=False)
                layer.close()
        done += 1
        progress(done, len(todo), f"page {i + 1}")

    # Pictures are made here one by one (MuPDF documents are not shared between threads), a few pages ahead of
    # Tesseract, which reads them in parallel; at most 2 × workers pictures are held at once.
    with ThreadPoolExecutor(max_workers=workers) as pool:
        for start in range(0, len(todo), workers * 2):
            batch = []
            for i in todo[start:start + workers * 2]:
                prep = prepare(i)
                batch.append((i, prep, pool.submit(read_picture, binary, prep["png"], language, psm, dpi, tessdata,
                                                   want_pdf)))
            for i, prep, future in batch:
                pdf, text = future.result()
                finish(i, prep, pdf, text)

    ordered = [(i + 1, texts.get(i, "")) for i in pages]
    outputs: list[T.Output] = []
    if want_pdf:
        outputs.append(T._pdf_out(out_dir, doc, T._name(inp.stem, "searchable"), garbage=3))
    doc.close()
    if output in ("txt", "pdf_txt"):
        outputs.append(T._file_out(out_dir, _text_file(ordered).encode("utf-8"), T._name(inp.stem, "text", ".txt")))
    if output == "docx":
        outputs.append(T._file_out(out_dir, _docx(ordered, inp.stem), T._name(inp.stem, "text", ".docx")))
    skipped = len(pages) - len(todo)
    chars = sum(len(t) for _, t in ordered)
    tail = f" {skipped} already had text." if skipped else ""
    tail += f" {straightened} page(s) straightened or turned upright." if straightened else ""
    if chars < 5 * len(todo):
        tail += " Very little text was found: check the language, or try the AI engine for hard pages."
    return Result(outputs, f"{len(todo)} page(s) read in {language.replace('+', ' + ')}.{tail}",
                  {"pages": len(todo), "language": language, "chars": chars})
