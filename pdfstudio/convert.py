"""Conversions: PDF → Word / Excel / PowerPoint, Office files → PDF, Text / Markdown / HTML → PDF.

- PDF → Word: "layout" keeps the page layout, tables and pictures (pdf2docx, when installed); "text" writes plain
  flowing paragraphs with headings and pictures, easy to edit (python-docx). Scans have no text: OCR them first.
- PDF → Excel: MuPDF finds the tables (ruled or by alignment); each table becomes a sheet, or all go on one sheet.
- PDF → PowerPoint: each page a full-slide picture, the page's text in the speaker notes (searchable).
- Office → PDF: LibreOffice in headless mode (each run with its own profile folder, so runs can overlap).
- Text / Markdown / HTML → PDF: MuPDF's own HTML layout (no browser, no LibreOffice); Bengali and other scripts
  are shaped and fall back to MuPDF's Noto fonts.
"""
from __future__ import annotations

import html
import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf

from . import capabilities as C
from . import tools as T
from .tools import Input, Progress, Result, ToolError

OFFICE_TIMEOUT = 300


# ---------------------------------------------------------------------------------------------- PDF → Word
def _empty_pages(doc: pymupdf.Document, pages: list[int]) -> int:
    return sum(1 for i in pages if not doc[i].get_text("text").strip())


def pdf2word(inputs: list[Input], p: dict, out_dir: Path, progress: Progress) -> Result:
    inp = inputs[0]
    doc = T.open_pdf(inp)
    pages = T.page_set(T._text(p, "pages", 4000), doc.page_count)
    empty = _empty_pages(doc, pages)
    mode = T._choice(p, "mode", ("layout", "text"), "layout")
    if empty == len(pages):
        doc.close()
        raise ToolError("These pages are scans (pictures, no text). Run “OCR: scans to text” first and choose Word "
                        "as the result, or make it searchable and convert that.")
    import importlib.util

    if importlib.util.find_spec("pdf2docx") is None:          # optional: the simple conversion needs only python-docx
        mode = "text"
    target = out_dir / "out-word.docx"
    out_dir.mkdir(parents=True, exist_ok=True)
    if mode == "layout":
        doc.close()
        import warnings

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            from pdf2docx import Converter

        progress(0, 2, "reading the layout")
        converter = Converter(str(inp.path))
        try:
            converter.convert(str(target), pages=pages)
        except Exception as exc:  # pdf2docx fails on some unusual files: fall back to flowing text
            converter.close()
            target.unlink(missing_ok=True)
            doc = T.open_pdf(inp)
            mode = "text"
            note = f" (the exact layout failed: {str(exc)[:120]}; plain text was used)"
        else:
            converter.close()
            note = ""
    else:
        note = ""
    if mode == "text":
        _simple_docx(doc, pages, target, progress)
        doc.close()
    name = T._name(inp.stem, "", ".docx")
    tail = f" {empty} page(s) had no text (scans): OCR them first." if empty else ""
    how = "with its layout, tables and pictures" if mode == "layout" else "as flowing text with pictures"
    return Result([T.Output(target, name)], f"{len(pages)} page(s) converted to Word {how}.{note}{tail}",
                  {"pages": len(pages), "mode": mode})


def _simple_docx(doc: pymupdf.Document, pages: list[int], target: Path, progress: Progress) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    out = Document()
    out.styles["Normal"].font.size = Pt(11)
    body_sizes = []
    for i in pages[:50]:
        for block in doc[i].get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    if span["text"].strip():
                        body_sizes.append(round(span["size"]))
    body = max(set(body_sizes), key=body_sizes.count) if body_sizes else 11
    width = out.sections[0].page_width - out.sections[0].left_margin - out.sections[0].right_margin
    for k, i in enumerate(pages):
        progress(k, len(pages), f"page {i + 1}")
        if k:
            out.add_page_break()
        page = doc[i]
        for block in page.get_text("dict", sort=True)["blocks"]:
            if block.get("type") == 1:                          # a picture
                try:
                    pix = pymupdf.Pixmap(block["image"]) if isinstance(block.get("image"), bytes) else None
                    data = block["image"] if pix is None else pix.tobytes("png")
                    ratio = (block["bbox"][2] - block["bbox"][0]) / max(page.rect.width, 1)
                    out.add_picture(io.BytesIO(data), width=int(width * max(0.15, min(1.0, ratio))))
                except Exception:  # an image format python-docx cannot take: skip it
                    pass
                continue
            spans = [s for line in block.get("lines", []) for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            size = max(s["size"] for s in spans)
            text = " ".join("".join(s["text"] for s in line["spans"]).strip() for line in block["lines"]).strip()
            text = re.sub(r"(\w)- (\w)", r"\1\2", text)
            if size >= body * 1.35 and len(text) < 200:
                out.add_heading(text, level=1 if size >= body * 1.7 else 2)
                continue
            para = out.add_paragraph()
            for line in block["lines"]:
                for s in line["spans"]:
                    run = para.add_run(s["text"])
                    run.bold = bool(s["flags"] & 16)
                    run.italic = bool(s["flags"] & 2)
                    if abs(s["size"] - body) > 1:
                        run.font.size = Pt(round(s["size"]))
                    if s["color"]:
                        run.font.color.rgb = RGBColor.from_string(f"{s['color'] & 0xFFFFFF:06X}")
                para.add_run(" ")
    out.save(target)


# ---------------------------------------------------------------------------------------------- PDF → Excel
def _cell(value) -> object:
    if value is None:
        return None
    text = re.sub(r"\s*\n\s*", " ", str(value)).strip()
    number = text.replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", number) and len(number) < 16 and not re.fullmatch(r"0\d+", number):
        return float(number) if "." in number else int(number)
    return text


def pdf2excel(inputs: list[Input], p: dict, out_dir: Path, progress: Progress) -> Result:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    inp = inputs[0]
    doc = T.open_pdf(inp)
    pages = T.page_set(T._text(p, "pages", 4000), doc.page_count)
    strategy = T._choice(p, "find", ("lines", "text"), "lines")
    layout = T._choice(p, "layout", ("sheets", "one"), "sheets")
    header = T._bool(p, "header", True)
    tables: list[tuple[int, list[list]]] = []
    for k, i in enumerate(pages):
        progress(k, len(pages), f"page {i + 1}")
        page = T._upright(doc[i])
        try:
            found = page.find_tables(strategy=strategy)
        except Exception:  # an odd page: no tables from it
            continue
        for table in found.tables:
            rows = table.extract()
            rows = [r for r in rows if any(c not in (None, "") for c in r)]
            if len(rows) >= 1 and max(len(r) for r in rows) >= 2:
                tables.append((i + 1, rows))
    if not tables and T._bool(p, "fallback", True):
        for i in pages:                                       # no table: every text line as a row, split on gaps
            words = doc[i].get_text("words", sort=True)
            lines: dict[tuple, list] = {}
            for w in words:
                lines.setdefault((w[5], w[6]), []).append(w)
            rows = []
            for ws in lines.values():
                cells: list[str] = []
                last = None
                for w in ws:                                  # a gap wider than ~2 spaces starts a new cell
                    if cells and last is not None and w[0] - last <= 14:
                        cells[-1] += " " + w[4]
                    else:
                        cells.append(w[4])
                    last = w[2]
                rows.append(cells)
            if rows:
                tables.append((i + 1, rows))
    doc.close()
    if not tables:
        raise ToolError("No tables were found. Try “find tables by text alignment”, or OCR a scan first.")
    book = Workbook()
    book.remove(book.active)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8ECF8")
    one = book.create_sheet("Tables") if layout == "one" else None
    row_at = 1
    for n, (page_no, rows) in enumerate(tables, 1):
        sheet = one or book.create_sheet(f"P{page_no} T{n}"[:31])
        start = row_at if one else 1
        if one:
            sheet.cell(start, 1, f"Table {n} · page {page_no}").font = Font(bold=True, size=12)
            start += 1
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = sheet.cell(start + r, c + 1, _cell(value))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if header and r == 0:
                    cell.font, cell.fill = bold, fill
        widths: dict[int, int] = {}
        for row in rows:
            for c, value in enumerate(row):
                widths[c] = max(widths.get(c, 6), min(60, len(str(value or "")) + 2))
        for c, w in widths.items():
            letter = get_column_letter(c + 1)
            sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width or 0, w)
        if not one:
            sheet.freeze_panes = "A2" if header else None
        row_at = start + len(rows) + 2
    target = out_dir / "out-excel.xlsx"
    out_dir.mkdir(parents=True, exist_ok=True)
    book.save(target)
    return Result([T.Output(target, T._name(inp.stem, "", ".xlsx"))],
                  f"{len(tables)} table(s) put into a spreadsheet.", {"tables": len(tables)})


# ---------------------------------------------------------------------------------------------- PDF → PowerPoint
def pdf2ppt(inputs: list[Input], p: dict, out_dir: Path, progress: Progress) -> Result:
    from pptx import Presentation
    from pptx.util import Emu

    inp = inputs[0]
    doc = T.open_pdf(inp)
    pages = T.page_list(T._text(p, "pages", 4000), doc.page_count)
    dpi = T._int(p, "dpi", 150, 72, 300)
    notes = T._bool(p, "notes", True)
    first = doc[pages[0]].rect
    deck = Presentation()
    deck.slide_width = Emu(int(first.width * 12700))
    deck.slide_height = Emu(int(first.height * 12700))
    blank = deck.slide_layouts[6]
    for k, i in enumerate(pages):
        progress(k, len(pages), f"page {i + 1}")
        page = doc[i]
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        data = pix.tobytes("jpg", jpg_quality=88) if pix.width * pix.height > 4_000_000 else pix.tobytes("png")
        slide = deck.slides.add_slide(blank)
        sw, sh = deck.slide_width, deck.slide_height
        scale = min(sw / (page.rect.width * 12700), sh / (page.rect.height * 12700))
        w, h = int(page.rect.width * 12700 * scale), int(page.rect.height * 12700 * scale)
        slide.shapes.add_picture(io.BytesIO(data), (sw - w) // 2, (sh - h) // 2, w, h)
        if notes:
            text = page.get_text("text").strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text[:20000]
    doc.close()
    target = out_dir / "out-slides.pptx"
    out_dir.mkdir(parents=True, exist_ok=True)
    deck.save(target)
    return Result([T.Output(target, T._name(inp.stem, "", ".pptx"))], f"{len(pages)} slide(s) made.",
                  {"slides": len(pages)})


# ---------------------------------------------------------------------------------------------- Office → PDF
def office2pdf(inputs: list[Input], p: dict, out_dir: Path, progress: Progress) -> Result:
    binary = C.libreoffice_path()
    if not binary:
        raise ToolError("Office → PDF uses LibreOffice (free), which is not installed on this computer. Install it "
                        f"with: {C.INSTALL['libreoffice'][C.SYSTEM]}")
    outputs: list[T.Output] = []
    one = T._bool(p, "merge") and len(inputs) > 1
    merged = pymupdf.open() if one else None
    for k, inp in enumerate(inputs):
        progress(k, len(inputs), inp.name)
        suffix = Path(inp.name).suffix.lower() or ".docx"
        with tempfile.TemporaryDirectory(prefix="phipdf-office-") as tmp:
            work = Path(tmp)
            src = work / f"input{suffix}"
            shutil.copyfile(inp.path, src)
            profile = (work / "profile").as_uri()
            cmd = [binary, f"-env:UserInstallation={profile}", "--headless", "--norestore", "--nologo",
                   "--convert-to", "pdf", "--outdir", str(work / "out"), str(src)]
            try:
                done = subprocess.run(cmd, capture_output=True, timeout=OFFICE_TIMEOUT,
                                      creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            except subprocess.TimeoutExpired as exc:
                raise ToolError(f"{inp.name} took too long to convert (over 5 minutes).") from exc
            result = work / "out" / "input.pdf"
            if done.returncode != 0 or not result.is_file():
                detail = (done.stderr or done.stdout).decode("utf-8", "replace").strip()[:200]
                raise ToolError(f"LibreOffice could not convert {inp.name}{': ' + detail if detail else '.'}")
            if merged is not None:
                with pymupdf.open(result) as part:
                    start = merged.page_count
                    merged.insert_pdf(part)
                    if part.page_count:
                        merged.set_toc(merged.get_toc() + [[1, Path(inp.name).stem, start + 1]])
            else:
                target = out_dir / f"out-{k:04d}.pdf"
                out_dir.mkdir(parents=True, exist_ok=True)
                shutil.move(str(result), target)
                outputs.append(T.Output(target, T._name(inp.stem, "")))
    if merged is not None:
        outputs.append(T._pdf_out(out_dir, merged, T._name(inputs[0].stem, f"+{len(inputs) - 1} more")))
        merged.close()
    return Result(outputs, f"{len(inputs)} file(s) converted to PDF.", {"files": len(inputs)})


# ---------------------------------------------------------------------------------------------- text → PDF
def _markdown(text: str) -> str:
    try:
        import markdown
    except ImportError:  # optional: paragraphs and line breaks only
        return "".join(f"<p>{html.escape(par).replace(chr(10), '<br>')}</p>" for par in re.split(r"\n\s*\n", text))
    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])


def _clean_html(source: str) -> str:
    """Only the body, without scripts, styles we cannot use, or remote resources (nothing is fetched)."""
    source = re.sub(r"(?is)<(script|style|iframe|object|noscript|template|svg)\b.*?</\1\s*>", "", source)
    source = re.sub(r"(?is)<(script|style|iframe|object|embed|link|meta|base)\b[^>]*>", "", source)
    body = re.search(r"(?is)<body[^>]*>(.*)</body>", source)
    source = body.group(1) if body else source
    return re.sub(r"(?i)\s(src|href)\s*=\s*(['\"])(?!#)[^'\"]*\2", "", source)


TEXT_CSS = """
body {{ font-family: {family}; font-size: {size}pt; line-height: 1.45; color: #111; }}
h1 {{ font-size: 1.7em; margin: 0 0 .5em; }} h2 {{ font-size: 1.35em; margin: 1em 0 .4em; }}
h3 {{ font-size: 1.15em; margin: .9em 0 .3em; }} p {{ margin: 0 0 .6em; text-align: {align}; }}
pre, code {{ font-family: monospace; font-size: .9em; }} pre {{ background: #f3f4f6; padding: 6pt; }}
table {{ border-collapse: collapse; margin: .4em 0 .8em; }} td, th {{ border: 1px solid #999; padding: 3pt 6pt; }}
th {{ background: #eef0f6; }} blockquote {{ margin: .4em 0 .6em 1em; padding-left: .6em; border-left: 3pt solid #ccc; }}
"""


def text2pdf(inputs: list[Input], p: dict, out_dir: Path, progress: Progress) -> Result:
    size = T._float(p, "size", 12, 6, 40)
    family = {"sans": "sans-serif", "serif": "serif", "mono": "monospace"}[T._choice(p, "font", tuple(T.FONTS), "serif")]
    pw, ph = T.PAPER[T._choice(p, "paper", tuple(T.PAPER), "A4")]
    if T._choice(p, "orientation", ("portrait", "landscape"), "portrait") == "landscape":
        pw, ph = ph, pw
    margin = T._float(p, "margin", 20, 5, 60) * T.MM
    align = "justify" if T._bool(p, "justify") else "left"
    css = TEXT_CSS.format(family=family, size=size, align=align)
    one = T._bool(p, "merge", True)
    box = pymupdf.Rect(margin, margin, pw - margin, ph - margin)

    def render(part: str) -> pymupdf.Document:
        story = pymupdf.Story(html=part, user_css=css)
        buffer = io.BytesIO()
        writer = pymupdf.DocumentWriter(buffer)
        more, count = 1, 0
        while more and count < T.MAX_PAGES:
            device = writer.begin_page(pymupdf.Rect(0, 0, pw, ph))
            more, _ = story.place(box)
            story.draw(device)
            writer.end_page()
            count += 1
        writer.close()
        return pymupdf.open("pdf", buffer.getvalue())

    outputs: list[T.Output] = []
    merged = pymupdf.open() if one else None
    pages_total = 0
    for k, inp in enumerate(inputs):
        progress(k, len(inputs), inp.name)
        raw = inp.path.read_bytes()[:20_000_000].decode("utf-8", "replace")
        suffix = Path(inp.name).suffix.lower()
        how = T._choice(p, "format", ("auto", "plain", "markdown"), "auto")
        if suffix in (".html", ".htm"):
            part = _clean_html(raw)
        elif suffix in (".md", ".markdown") or how == "markdown":
            part = _markdown(raw)
        else:
            part = "".join(f"<p>{html.escape(par).replace(chr(10), '<br>')}</p>"
                           for par in re.split(r"\n\s*\n", raw.replace("\r\n", "\n")) if par.strip())
        doc = render(part or "<p></p>")
        pages_total += doc.page_count
        if merged is not None:
            start = merged.page_count
            merged.insert_pdf(doc)
            if len(inputs) > 1:
                merged.set_toc(merged.get_toc() + [[1, inp.stem, start + 1]])
        else:
            outputs.append(T._pdf_out(out_dir, doc, T._name(inp.stem, "")))
        doc.close()
    if merged is not None:
        name = T._name(inputs[0].stem, "" if len(inputs) == 1 else f"+{len(inputs) - 1} more")
        outputs.append(T._pdf_out(out_dir, merged, name))
        merged.close()
    return Result(outputs, f"{len(inputs)} file(s) set as PDF: {pages_total} page(s).", {"pages": pages_total})
